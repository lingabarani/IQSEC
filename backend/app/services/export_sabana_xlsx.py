"""
Sábana Matrix Excel Exporter Service
Generates the comprehensive, color-coded RFP Compliance Matrix (.xlsx)
matching Mexican public tender and enterprise procurement standards.
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List
from backend.app.db.models.requirement import RFPRequirement, ComplianceStatus


class SabanaXlsxExporter:
    """
    Exports RFP Requirements & AI Evaluations to an executive Excel Sábana Matrix.
    """

    # Brand Colors
    HEADER_FILL = PatternFill(start_color="0A192F", end_color="0A192F", fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    STATUS_FILLS = {
        ComplianceStatus.COMPLIES: PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
        ComplianceStatus.COMPLIES_WITH_EXCEPTION: PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
        ComplianceStatus.DOES_NOT_COMPLY: PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
        ComplianceStatus.NOT_ENOUGH_EVIDENCE: PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid"),
        ComplianceStatus.NOT_EVALUATED: PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
    }

    STATUS_FONTS = {
        ComplianceStatus.COMPLIES: Font(name="Calibri", size=10, bold=True, color="155724"),
        ComplianceStatus.COMPLIES_WITH_EXCEPTION: Font(name="Calibri", size=10, bold=True, color="856404"),
        ComplianceStatus.DOES_NOT_COMPLY: Font(name="Calibri", size=10, bold=True, color="721C24"),
        ComplianceStatus.NOT_ENOUGH_EVIDENCE: Font(name="Calibri", size=10, bold=True, color="383D41"),
        ComplianceStatus.NOT_EVALUATED: Font(name="Calibri", size=10, color="000000"),
    }

    THIN_BORDER = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )

    def generate_sabana_workbook(
        self,
        rfp_title: str,
        tender_number: str,
        requirements: List[RFPRequirement],
        output_filepath: str
    ) -> str:
        """
        Creates and saves the formatted Sábana Excel Matrix.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Matriz de Cumplimiento (Sábana)"

        # Title Block
        ws.merge_cells("A1:K1")
        title_cell = ws["A1"]
        title_cell.value = f"IQSEC S.A. DE C.V. - MATRIZ DE CUMPLIMIENTO TÉCNICO Y ECONÓMICO (SÁBANA)"
        title_cell.font = Font(name="Calibri", size=14, bold=True, color="0A192F")
        title_cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A2:K2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Licitación: {tender_number} | {rfp_title}"
        sub_cell.font = Font(name="Calibri", size=11, italic=True, color="555555")

        # Headers
        headers = [
            "No.",
            "Pág.",
            "Numeral / Cláusula",
            "Sección del Anexo",
            "Pilar IQSEC",
            "Texto del Requerimiento",
            "Modificado en Junta?",
            "Dictamen",
            "Propuesta Técnica IQSEC",
            "Cita Documental",
            "Confianza"
        ]

        header_row = 4
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=header)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.row_dimensions[header_row].height = 28

        # Populate Rows
        for r_idx, req in enumerate(requirements, 1):
            row_num = header_row + r_idx

            # Build citation text
            citations_str = "N/A"
            if req.exact_citations:
                c_items = []
                for c in req.exact_citations:
                    c_items.append(f"{c.get('document_title', 'Doc')} (Pág. {c.get('page_number', 1)})")
                citations_str = "\n".join(c_items)

            mod_str = "NO"
            if req.is_modified_by_addendum:
                mod_str = f"SÍ ({req.addendum_question_num or 'Junta'} Pág. {req.addendum_page_num or 1})"

            status_val = req.compliance_status.value if hasattr(req.compliance_status, "value") else str(req.compliance_status)

            values = [
                r_idx,
                req.page_number,
                req.requirement_code,
                req.section_title[:45],
                req.iqsec_pillar.value if hasattr(req.iqsec_pillar, "value") else str(req.iqsec_pillar),
                req.effective_text,
                mod_str,
                status_val,
                req.technical_response or "Pendiente de evaluación",
                citations_str,
                f"{round(req.confidence_score * 100, 1)}%"
            ]

            for c_idx, val in enumerate(values, 1):
                c = ws.cell(row=row_num, column=c_idx, value=val)
                c.border = self.THIN_BORDER
                c.font = Font(name="Calibri", size=9)
                c.alignment = Alignment(vertical="top", wrap_text=(c_idx in [6, 9, 10]))

                # Format status cell
                if c_idx == 8:
                    c.fill = self.STATUS_FILLS.get(req.compliance_status, self.STATUS_FILLS[ComplianceStatus.NOT_EVALUATED])
                    c.font = self.STATUS_FONTS.get(req.compliance_status, self.STATUS_FONTS[ComplianceStatus.NOT_EVALUATED])
                    c.alignment = Alignment(horizontal="center", vertical="center")

                if c_idx in [1, 2, 7, 11]:
                    c.alignment = Alignment(horizontal="center", vertical="top")

            ws.row_dimensions[row_num].height = 45

        # Column widths
        col_widths = {
            1: 6,   # No.
            2: 7,   # Pág.
            3: 18,  # Numeral
            4: 24,  # Sección
            5: 18,  # Pilar
            6: 45,  # Requerimiento
            7: 18,  # Modificado
            8: 16,  # Dictamen
            9: 50,  # Propuesta
            10: 28, # Citas
            11: 12  # Confianza
        }
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # Freeze Panes
        ws.freeze_panes = "A5"

        os.makedirs(os.path.dirname(output_filepath), exist_ok=True)
        wb.save(output_filepath)
        return output_filepath


sabana_exporter = SabanaXlsxExporter()

