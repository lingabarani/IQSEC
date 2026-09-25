"""
Technical Proposal Word Document (.docx) Exporter Service
Generates the formal Technical Proposal document complying with Mexican public and private procurement standards.
"""
import os
from typing import List
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from backend.app.db.models.requirement import RFPRequirement, ComplianceStatus


class TechnicalDocxExporter:
    """
    Generates professional Technical Proposal Word Documents (.docx).
    """

    def generate_technical_proposal(
        self,
        rfp_title: str,
        tender_number: str,
        customer_name: str,
        requirements: List[RFPRequirement],
        output_filepath: str
    ) -> str:
        doc = docx.Document()

        # Set Margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # Title Page
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_company = title_p.add_run("IQSEC S.A. DE C.V.\n")
        run_company.font.name = "Calibri"
        run_company.font.size = Pt(22)
        run_company.font.bold = True
        run_company.font.color.rgb = RGBColor(10, 25, 47)

        run_sub = title_p.add_run("LÍDER EN CIBERSEGURIDAD, IDENTIDAD DIGITAL Y SERVICIOS SOC GESTIONADOS\n\n\n")
        run_sub.font.name = "Calibri"
        run_sub.font.size = Pt(11)
        run_sub.font.color.rgb = RGBColor(0, 180, 216)

        run_title = title_p.add_run(f"PROPUESTA TÉCNICA DETALLADA\n{rfp_title}\n\n")
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(16)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(10, 25, 47)

        meta_p = doc.add_paragraph()
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_run = meta_p.add_run(
            f"Licitación / Procedimiento: {tender_number}\n"
            f"Cliente / Convocante: {customer_name}\n"
            f"Fecha de Presentación: Septiembre 2026\n"
            f"Versión: 1.0 Final Aprobada"
        )
        meta_run.font.name = "Calibri"
        meta_run.font.size = Pt(11)
        meta_run.font.italic = True

        doc.add_page_break()

        # Section 1: Executive Summary
        h1 = doc.add_heading("1. RESUMEN EJECUTIVO Y PERFIL DE IQSEC", level=1)
        h1.style.font.color.rgb = RGBColor(10, 25, 47)

        p_exec = doc.add_paragraph()
        p_exec.add_run(
            "IQSEC S.A. de C.V. es una empresa 100% mexicana con más de 18 años de experiencia en el diseño, "
            "implementación y operación de servicios avanzados de ciberseguridad, respuesta a incidentes, "
            "identidad digital y Centros de Operaciones de Seguridad (SOC) certificados bajo ISO/IEC 27001:2022.\n\n"
            "La presente propuesta técnica integra soluciones robustas de última generación, garantizando "
            "la continuidad operativa, confidencialidad, integridad y estricto apego a los marcos regulatorios "
            "aplicables (CNBV, PCI-DSS, NIST)."
        )

        # Section 2: Technical Compliance Table
        h2 = doc.add_heading("2. MATRIZ DE CUMPLIMIENTO TÉCNICO Y PROPUESTA OPERATIVA", level=1)
        h2.style.font.color.rgb = RGBColor(10, 25, 47)

        # Table
        table = doc.add_table(rows=1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        headers = ["No.", "Numeral", "Especificación Requerida", "Dictamen", "Propuesta Técnica y Evidencia"]
        widths = [Inches(0.5), Inches(1.1), Inches(2.2), Inches(1.1), Inches(2.6)]

        # Header formatting
        hdr_cells = table.rows[0].cells
        for idx, (header_text, w) in enumerate(zip(headers, widths)):
            hdr_cells[idx].text = header_text
            hdr_cells[idx].width = w
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0A192F"/>')
            hdr_cells[idx]._tc.get_or_add_tcPr().append(shading)
            for p in hdr_cells[idx].paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
                    r.font.size = Pt(9.5)

        for r_idx, req in enumerate(requirements, 1):
            row_cells = table.add_row().cells
            for idx, w in enumerate(widths):
                row_cells[idx].width = w

            row_cells[0].text = str(r_idx)
            row_cells[1].text = req.requirement_code
            row_cells[2].text = req.effective_text[:200]
            
            # Status
            status_text = req.compliance_status.value if hasattr(req.compliance_status, "value") else str(req.compliance_status)
            row_cells[3].text = status_text
            
            # Response + citations
            c_text = req.technical_response or "Cumple conforme a especificaciones."
            if req.exact_citations:
                c_text += "\n\nEvidencia Documental:\n" + "\n".join(
                    [f"• {c.get('document_title', '')} (Pág. {c.get('page_number', 1)})" for c in req.exact_citations]
                )
            row_cells[4].text = c_text

            for cell in row_cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8.5)
                        r.font.name = "Calibri"

        os.makedirs(os.path.dirname(output_filepath), exist_ok=True)
        doc.save(output_filepath)
        return output_filepath


technical_docx_exporter = TechnicalDocxExporter()

