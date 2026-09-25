"""
Executive Presentation Slide Deck (.pptx) Exporter Service
Generates high-impact executive presentation decks with IQSEC branding.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from backend.app.db.models.proposal import Proposal


class ExecutivePptxExporter:
    """
    Generates professional 16:9 widescreen Executive Presentations (.pptx).
    """

    COLOR_DARK = RGBColor(10, 25, 47)      # Deep Navy #0A192F
    COLOR_CYAN = RGBColor(0, 180, 216)     # Cyan #00B4D8
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_GRAY = RGBColor(240, 242, 245)
    COLOR_TEXT_DARK = RGBColor(33, 37, 41)

    def generate_executive_deck(
        self,
        proposal: Proposal,
        tender_number: str,
        customer_name: str,
        output_filepath: str
    ) -> str:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        blank_layout = prs.slide_layouts[6]

        # ----------------------------------------------------
        # Slide 1: Cover Slide
        # ----------------------------------------------------
        slide1 = prs.slides.add_slide(blank_layout)
        
        # Background
        bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = self.COLOR_DARK
        bg.line.fill.background()

        # Title text
        txBox = slide1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(3.5))
        tf = txBox.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = "IQSEC S.A. DE C.V."
        p1.font.name = "Arial"
        p1.font.size = Pt(28)
        p1.font.bold = True
        p1.font.color.rgb = self.COLOR_CYAN

        p2 = tf.add_paragraph()
        p2.text = f"Propuesta Técnica y Ejecutiva\n{proposal.title}"
        p2.font.name = "Arial"
        p2.font.size = Pt(36)
        p2.font.bold = True
        p2.font.color.rgb = self.COLOR_WHITE

        p3 = tf.add_paragraph()
        p3.text = f"Licitación: {tender_number} | Cliente: {customer_name}"
        p3.font.name = "Arial"
        p3.font.size = Pt(16)
        p3.font.color.rgb = RGBColor(180, 200, 220)

        # ----------------------------------------------------
        # Slide 2: Compliance Metrics Dashboard
        # ----------------------------------------------------
        slide2 = prs.slides.add_slide(blank_layout)
        self._add_slide_header(slide2, "RESUMEN EJECUTIVO DE CUMPLIMIENTO", "Evaluación cuantitativa de bases de licitación")

        # Metric Cards
        metrics = [
            ("Requerimientos Totales", str(proposal.total_requirements), self.COLOR_DARK),
            ("Cumplimiento Favorable", f"{proposal.compliant_count} ({proposal.overall_compliance_rate}%)", RGBColor(40, 167, 69)),
            ("Cumple con Excepción", str(proposal.exception_count), RGBColor(255, 193, 7)),
            ("No Cumple", str(proposal.non_compliant_count), RGBColor(220, 53, 69)),
        ]

        card_width = Inches(2.6)
        card_height = Inches(3.2)
        start_left = Inches(1.0)
        gap = Inches(0.4)

        for i, (label, val, col) in enumerate(metrics):
            card_left = start_left + i * (card_width + gap)
            card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_left, Inches(2.2), card_width, card_height)
            card.fill.solid()
            card.fill.fore_color.rgb = self.COLOR_GRAY
            card.line.color.rgb = col
            card.line.width = Pt(2)

            tf_card = card.text_frame
            tf_card.word_wrap = True
            
            p_val = tf_card.paragraphs[0]
            p_val.text = f"\n\n{val}"
            p_val.font.size = Pt(32)
            p_val.font.bold = True
            p_val.font.color.rgb = col
            p_val.alignment = PP_ALIGN.CENTER

            p_lbl = tf_card.add_paragraph()
            p_lbl.text = f"\n{label}"
            p_lbl.font.size = Pt(14)
            p_lbl.font.color.rgb = self.COLOR_TEXT_DARK
            p_lbl.alignment = PP_ALIGN.CENTER

        # ----------------------------------------------------
        # Slide 3: SLA & Service Commitments
        # ----------------------------------------------------
        slide3 = prs.slides.add_slide(blank_layout)
        self._add_slide_header(slide3, "NIVELES DE SERVICIO Y OPERACIÓN SOC", "Garantías contractuales y esquema de entrega 24/7/365")

        box_sla = slide3.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(4.5))
        tf_sla = box_sla.text_frame
        tf_sla.word_wrap = True

        sla_items = [
            ("Centro de Operaciones de Seguridad (SOC 24/7)", "Operación ininterrumpida 365 días al año certificada bajo ISO/IEC 27001:2022."),
            ("Tiempo de Notificación Crítica (SLA)", "Notificación y triaje inicial de incidentes de severidad crítica en menos de 15 minutos."),
            ("Disponibilidad de Plataforma", "Garantía de disponibilidad del servicio del 99.95% con penalizaciones pactadas."),
            ("Célula de Ingeniería Especializada", "Asignación de Ingenieros de Seguridad certificados (CISSP, CISM, CEH, OSCP).")
        ]

        for title, desc in sla_items:
            p_t = tf_sla.add_paragraph()
            p_t.text = f"• {title}"
            p_t.font.bold = True
            p_t.font.size = Pt(16)
            p_t.font.color.rgb = self.COLOR_DARK

            p_d = tf_sla.add_paragraph()
            p_d.text = f"   {desc}\n"
            p_d.font.size = Pt(13)
            p_d.font.color.rgb = self.COLOR_TEXT_DARK

        os.makedirs(os.path.dirname(output_filepath), exist_ok=True)
        prs.save(output_filepath)
        return output_filepath

    def _add_slide_header(self, slide, title: str, subtitle: str):
        tx = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.333), Inches(1.2))
        tf = tx.text_frame
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = self.COLOR_DARK

        p_s = tf.add_paragraph()
        p_s.text = subtitle
        p_s.font.size = Pt(12)
        p_s.font.italic = True
        p_s.font.color.rgb = self.COLOR_CYAN


executive_pptx_exporter = ExecutivePptxExporter()

